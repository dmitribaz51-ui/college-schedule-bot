"""Анализ РЕАЛЬНОЙ структуры Excel-файла расписания.

Запуск:
    python -m tools.analyze_excel data/excel/2026-09-08_schedule.xlsx
    python -m tools.analyze_excel data/excel/file.xlsx --rows 40
"""
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook

from app.parser.excel_parser import (
    GROUP_RE, _find_groups_in_row, _find_lesson_number, _sheet_grid,
)


def analyze(path: Path, rows_limit: int) -> None:
    workbook = load_workbook(path, data_only=True)
    print(f"\n=== ФАЙЛ: {path.name} ===")
    print(f"Листы: {workbook.sheetnames}\n")

    for ws in workbook.worksheets:
        print("-" * 70)
        print(f"ЛИСТ: '{ws.title}'  строк={ws.max_row} столбцов={ws.max_column}")
        merged = list(ws.merged_cells.ranges)
        print(f"Объединённых диапазонов: {len(merged)}")
        for rng in merged[:15]:
            print(f"   merge {rng}")
        if len(merged) > 15:
            print(f"   ... ещё {len(merged) - 15}")

        grid = _sheet_grid(ws)
        print("\nСтроки (после разворота объединённых ячеек):")
        for r, row in enumerate(grid[:rows_limit], start=1):
            cells = [(c + 1, v) for c, v in enumerate(row) if v]
            if not cells:
                continue
            found = _find_groups_in_row(row)
            mark = f"  <== ГРУППЫ: {[g for _, g in found]}" if found else ""
            number = _find_lesson_number(row, found[0][0] if found else 5)
            if number is not None:
                mark += f"  <== НОМЕР ПАРЫ: {number}"
            preview = " | ".join(f"C{c}:{v[:35]}" for c, v in cells[:10])
            print(f" R{r:>3}: {preview}{mark}")
        print()

    workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="Анализ структуры Excel-расписания")
    parser.add_argument("path", help="путь к .xlsx файлу")
    parser.add_argument("--rows", type=int, default=30, help="сколько строк показать")
    args = parser.parse_args()

    path = Path(args.path)
    if not path.exists():
        raise SystemExit(f"Файл не найден: {path}")
    analyze(path, args.rows)


if __name__ == "__main__":
    main()