from pathlib import Path
from openpyxl import load_workbook

files = list(Path("data").rglob("*.xlsx"))

if not files:
    print("Excel-файлы не найдены в папке data")
    raise SystemExit

for path in files:
    print("\n" + "=" * 100)
    print(f"ФАЙЛ: {path}")
    print("=" * 100)

    try:
        workbook = load_workbook(path, data_only=True)

        for worksheet in workbook.worksheets:
            print(f"\n--- ЛИСТ: {worksheet.title!r} ---")
            print(
                f"Строк: {worksheet.max_row}, "
                f"столбцов: {worksheet.max_column}"
            )

            for row_number, row in enumerate(
                worksheet.iter_rows(values_only=True), start=1
            ):
                values = [
                    str(value).replace("\n", " ").strip()
                    for value in row
                    if value is not None and str(value).strip()
                ]

                if values:
                    print(f"{row_number:04d}: " + " | ".join(values))

        workbook.close()

    except Exception as error:
        print(f"ОШИБКА: {error}")
