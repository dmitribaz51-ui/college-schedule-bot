"""Диагностика Excel-файла для отладки парсинга кабинетов"""
import sys
from pathlib import Path
from openpyxl import load_workbook

# Исправление кодировки для Windows консоли
if sys.platform == 'win32':
    import io
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

def analyze_excel(file_path: str):
    print(f"Анализ файла: {file_path}\n")
    
    wb = load_workbook(file_path, data_only=True)
    ws = wb.active
    
    print(f"Лист: {ws.title}")
    print(f"Размеры: {ws.max_row} строк x {ws.max_column} колонок\n")
    
    print("=" * 80)
    print("ОБРАЗЕЦ ДАННЫХ (первые 25 строк)")
    print("=" * 80)
    
    for row_idx in range(1, min(26, ws.max_row + 1)):
        row_data = []
        for col_idx in range(1, min(15, ws.max_column + 1)):
            cell = ws.cell(row=row_idx, column=col_idx)
            value = cell.value
            if value:
                row_data.append(f"[{col_idx}]: {repr(value)}")
        
        if row_data:
            print(f"\nСтрока {row_idx}:")
            for item in row_data:
                print(f"  {item}")
    
    print("\n" + "=" * 80)
    print("ПОИСК ГРУППЫ ТД-24-9")
    print("=" * 80)
    
    target_group = "ТД-24-9"
    group_row = None
    group_col = None
    
    for row_idx in range(1, ws.max_row + 1):
        for col_idx in range(1, ws.max_column + 1):
            cell_value = ws.cell(row=row_idx, column=col_idx).value
            if cell_value and target_group in str(cell_value):
                print(f"✓ Найдена группа в строке {row_idx}, колонке {col_idx}: {repr(cell_value)}")
                group_row = row_idx
                group_col = col_idx
                break
        if group_row:
            break
    
    if not group_row:
        print(f"✗ Группа {target_group} не найдена!")
        wb.close()
        return
    
    print(f"\n{'=' * 80}")
    print(f"ДАННЫЕ ДЛЯ ГРУППЫ {target_group} (колонка {group_col})")
    print("=" * 80)
    
    # Показываем данные в колонке группы, начиная со строки после заголовка
    for row_idx in range(group_row + 1, min(group_row + 20, ws.max_row + 1)):
        # Номер пары (обычно в колонке 1 или 2)
        lesson_num = None
        for check_col in range(1, min(group_col, 4)):
            val = ws.cell(row=row_idx, column=check_col).value
            if val:
                val_str = str(val).strip()
                # Проверяем римские цифры или арабские
                if val_str in ['I', 'II', 'III', 'IV', 'V', 'VI', 'VII', 'VIII'] or val_str.isdigit():
                    lesson_num = val_str
                    break
        
        # Данные ячейки группы
        group_cell = ws.cell(row=row_idx, column=group_col).value
        
        if lesson_num or group_cell:
            print(f"\nСтрока {row_idx}:")
            if lesson_num:
                print(f"  Пара: {lesson_num}")
            if group_cell:
                print(f"  Содержимое: {repr(group_cell)}")
                
                # Пытаемся найти кабинет в тексте
                text = str(group_cell)
                if '/' in text:
                    parts = text.split('/')
                    print(f"  ↳ Разделение по '/': {parts}")
                    if len(parts) >= 2:
                        possible_room = parts[-1].strip()
                        print(f"  ↳ Потенциальный кабинет: {repr(possible_room)}")
    
    wb.close()

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Использование: python debug_excel.py <путь_к_файлу.xlsx>")
        sys.exit(1)
    
    file_path = sys.argv[1]
    if not Path(file_path).exists():
        print(f"Файл не найден: {file_path}")
        sys.exit(1)
    
    analyze_excel(file_path)
