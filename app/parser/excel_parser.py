"""ШАГ 14-17: разбор Excel-файла расписания.

ВАЖНО: структура файлов ПКПС заранее не задана, поэтому парсер работает
по признакам, а не по фиксированным адресам ячеек:
  1) объединённые ячейки «разворачиваются» в обычную сетку;
  2) строка считается заголовком, если в ней >= 2 названий групп (ТД-26-9);
  3) под заголовком ищутся номера пар в левых колонках;
  4) текст пары разбирается на предмет / преподавателя / кабинет.

После анализа реального файла (tools/analyze_excel.py) правки нужны ТОЛЬКО здесь.
"""
from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
import xlrd

from app.utils import extract_date, normalize_group

logger = logging.getLogger(__name__)

GROUP_RE = re.compile(
    r"[А-ЯЁA-Z]{1,6}\s*[-–—]\s*\d{1,2}"
    r"\s*[-–—]\s*\d{1,2}"
    r"(?:[А-ЯЁA-Z](?:\s*[-–—]\s*\d+)?"
    r"|\s*[-–—]\s*\d+)?",
    re.IGNORECASE,
)
COURSE_RE = re.compile(r"([1-4])\s*(?:[-–—]\s*[1-4]\s*)?курс", re.IGNORECASE)
ROMAN_COURSE_RE = re.compile(r"\b([IVX]{1,3})\s*курс\b", re.IGNORECASE)
LESSON_NUM_RE = re.compile(r"^\s*([0-8])\s*(?:пара|пар[аы]?|п\.)?\s*$", re.IGNORECASE)

ROMAN_TO_NUM = {
    "I": 1,
    "II": 2,
    "III": 3,
    "IV": 4,
    "V": 5,
    "VI": 6,
    "VII": 7,
    "VIII": 8,
}

LESSON_CELL_RE = re.compile(
    r"^\s*(?:(?P<roman>[IVX]+)|(?P<arabic>\d+))?"
    r"\s*(?P<time>\d{1,2}[:.]\d{2}(?::\d{2})?)?"
    r"\s*(?:пара|пар[аы]?|п\.)?\s*$",
    re.IGNORECASE,
)


def parse_lesson_number(value) -> int | None:
    if value is None:
        return None

    text = str(value).replace("l", "I").replace("L", "I").strip()
    match = LESSON_CELL_RE.match(text)

    if not match:
        return None

    # Ячейка с чистым временем («08:30:00», «14:10:00») — это не номер пары.
    # Римская цифра с временем («I 09:15») — номер пары.
    if match.group("time") and not match.group("roman"):
        return None

    if match.group("roman"):
        return ROMAN_TO_NUM.get(match.group("roman").upper())

    if match.group("arabic"):
        number = int(match.group("arabic"))
        return number if 1 <= number <= 8 else None

    return None


TEACHER_RE = re.compile(
    r"\b[А-ЯЁ][а-яё]+(?:-[А-ЯЁ][а-яё]+)?\s+[А-ЯЁ](?:\s*\.\s*[А-ЯЁ]?\.?|\s*[А-ЯЁ]\.?(?![а-яёА-ЯЁ]))"
)
ROOM_KEYWORD_RE = re.compile(
    r"(?:каб\.?|ауд\.?|кабинет|аудитория)\s*№?\s*([\w/\-]{1,10})", re.IGNORECASE
)
ROOM_SPECIAL_RE = re.compile(
    r"((?:(?:верхний|нижний|малый|большой)\s+)?"
    r"(?:спортивный\s+зал|спорт\s*зал|тренаж[её]рный\s+зал)"
    r"|с/з|акт\s*зал|библиотека)",
    re.IGNORECASE,
)
ROOM_PLAIN_RE = re.compile(r"(?<![\w/])(\d{1,3}[а-яА-Я]?)(?![\w/])")
NOTES_RE = re.compile(r"\(([^)]{2,60})\)")
MAX_HEADER_CELL_LEN = 40


@dataclass
class ParsedLesson:
    number: int
    subject: str | None = None
    teacher: str | None = None
    room: str | None = None
    notes: str | None = None

    def to_dict(self) -> dict:
        return {
            "number": self.number,
            "subject": self.subject or None,
            "teacher": self.teacher or None,
            "room": self.room or None,
            "notes": self.notes or None,
        }


@dataclass
class GroupSchedule:
    group: str
    schedule_date: date | None
    course: int | None = None
    lessons: list[ParsedLesson] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {
            "date": self.schedule_date.isoformat() if self.schedule_date else None,
            "group": self.group,
            "course": self.course,
            "lessons": [l.to_dict() for l in sorted(self.lessons, key=lambda x: x.number)],
        }


@dataclass
class _GroupColumn:
    group: str
    start: int
    end: int


# --------------------------------------------------------------------- утилиты
def _clean(text: str) -> str:
    return " ".join(str(text).replace("\xa0", " ").split())


def _sheet_grid(ws) -> list[list[str]]:
    """Матрица строк с «развёрнутыми» объединёнными ячейками."""
    max_row, max_col = ws.max_row or 0, ws.max_column or 0
    if not max_row or not max_col:
        return []
    grid = [["" for _ in range(max_col)] for _ in range(max_row)]

    for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
        for cell in row:
            value = cell.value
            if value is None:
                continue
            if isinstance(value, (datetime, date)):
                text = value.strftime("%d.%m.%Y")
            elif isinstance(value, float) and value.is_integer():
                text = str(int(value))
            else:
                text = _clean(value)
            if text:
                grid[cell.row - 1][cell.column - 1] = text

    for rng in ws.merged_cells.ranges:
        r0, c0 = rng.min_row - 1, rng.min_col - 1
        if r0 >= max_row or c0 >= max_col:
            continue
        value = grid[r0][c0]
        if not value:
            continue
        for r in range(r0, min(rng.max_row, max_row)):
            for c in range(c0, min(rng.max_col, max_col)):
                grid[r][c] = value
    return grid


def _find_groups_in_row(row: list[str]) -> list[tuple[int, str]]:
    found: list[tuple[int, str]] = []
    for col, text in enumerate(row):
        if not text or len(text) > MAX_HEADER_CELL_LEN:
            continue
        match = GROUP_RE.search(text)
        if not match:
            continue
        group = normalize_group(match.group(0))
        if group:
            found.append((col, group))
    return found


def _build_columns(found: list[tuple[int, str]], row_len: int) -> list[_GroupColumn]:
    """Из найденных ячеек делает диапазоны колонок для каждой группы."""
    spans: dict[str, list[int]] = {}
    for col, group in found:
        spans.setdefault(group, []).append(col)

    columns = [
        _GroupColumn(group=g, start=min(cols), end=max(cols))
        for g, cols in spans.items()
    ]
    columns.sort(key=lambda c: c.start)
    for i, column in enumerate(columns):
        next_start = columns[i + 1].start if i + 1 < len(columns) else row_len
        column.end = max(column.start, min(column.end, next_start - 1))
    return columns


def _find_lesson_number(row: list[str], limit_col: int) -> int | None:
    for col in range(0, min(limit_col, len(row))):
        number = parse_lesson_number(row[col])
        if number is not None:
            return number
    return None


def _cell_texts(row: list[str], start: int, end: int, group: str) -> list[str]:
    texts: list[str] = []
    for col in range(start, min(end, len(row) - 1) + 1):
        text = row[col]
        if not text or normalize_group(text) == group:
            continue
        if text not in texts:          # объединённые ячейки дают дубли
            texts.append(text)
    return texts


def split_lesson_text(text: str) -> ParsedLesson:
    """Разбирает текст пары на предмет / преподавателя / кабинет / примечание."""
    original = _clean(text)
    rest = original

    # «с 10.30 ВПР …» — пара начинается не по звонку. Время уходит в примечание,
    # чтобы не торчать в названии предмета.
    start_note = None
    start_match = re.match(r"^\s*с\s*(\d{1,2})[.](\d{2})\s+", rest, re.IGNORECASE)
    if start_match:
        start_note = f"с {int(start_match.group(1)):02d}:{start_match.group(2)}"
        rest = rest[start_match.end():]

    teacher = None
    match = TEACHER_RE.search(rest)
    if match:
        teacher = _clean(match.group(0)).replace(" .", ".").replace(". ", ".")
        rest = rest[: match.start()] + " " + rest[match.end():]

    room = None

    # Основной формат: преподаватель / кабинет.
    # Берём последний слэш с номером: «каб.410/ 406» — это кабинет 406, не 410.
    slash_room = None
    for slash_room in re.finditer(
        r"/\s*([0-9]{1,3}[А-Яа-яA-Za-z]?)\s*",
        rest,
    ):
        pass

    if slash_room:
        room = slash_room.group(1).strip()
        rest = rest[: slash_room.start()] + " " + rest[slash_room.end():]
    else:
        match = ROOM_KEYWORD_RE.search(rest)
        if match:
            room = _clean(match.group(1))
            rest = rest[: match.start()] + " " + rest[match.end():]
        else:
            match = ROOM_SPECIAL_RE.search(rest)
            if match:
                room = _clean(match.group(1))
                rest = rest[: match.start()] + " " + rest[match.end():]

    notes = start_note
    match = NOTES_RE.search(rest)
    if match:
        extra = _clean(match.group(1))
        notes = f"{notes}; {extra}" if notes else extra
        rest = rest[: match.start()] + " " + rest[match.end():]

    # «каб.410/ 406»: номер после слэша уже стал кабинетом, «каб.410» в предмете лишний
    if room:
        rest = re.sub(
            r"(?:каб\.?|ауд\.?|кабинет|аудитория)\s*№?\s*[\w/\-]*",
            " ", rest, flags=re.IGNORECASE,
        )

    subject = _clean(rest).strip(" /-–—,;.")
    if not subject:
        subject = original
    return ParsedLesson(
        number=0, subject=subject or None, teacher=teacher, room=room, notes=notes
    )


def _make_lesson(number: int, text: str) -> ParsedLesson:
    lesson = split_lesson_text(text)
    lesson.number = number
    return lesson


MIDDAY_CLASS_HOUR = 90  # условный номер: классный час середины дня (14:10-14:45)
CLASS_HOUR_TIME_PREFIX_RE = re.compile(r"^\s*с\s*\d{1,2}[.]\d{2}\s*", re.IGNORECASE)


def _split_class_hour_cell(text: str) -> list[tuple[bool, str]]:
    """Разделяет ячейку, где классный час склеен с парой через '/'.

    Порядок частей любой:
      'Обществознание Пепеляева А.А. каб.212/ с 14.10 Классный час Боровкова Ю.Д. каб.216'
        -> [(False, 'Обществознание …'), (True, 'с 14.10 Классный час …')]
      'с 14.10 Классный час Балмышева С.И. каб.302/ Технология … Горбунова Г.В. Каб.308'
        -> [(True, 'с 14.10 Классный час …'), (False, 'Технология …')]
    Короткий сегмент без преподавателя ('Спортзал', 'каб. N') — продолжение
    предыдущей части, а не отдельная пара.
    """
    cleaned = _clean(text)
    if "классный час" not in cleaned.lower():
        return [(False, cleaned)]

    segments = [s.strip() for s in cleaned.split("/") if s.strip()]
    if len(segments) <= 1:
        return [(True, cleaned)]

    results: list[tuple[bool, str]] = []
    for seg in segments:
        if "классный час" in seg.lower():
            results.append((True, seg))
            continue
        prev = results[-1] if results else None
        is_room_like = not TEACHER_RE.search(seg) and (
            bool(ROOM_SPECIAL_RE.search(seg)) or len(seg) <= 10
        )
        if prev is not None and is_room_like:
            kind, prev_text = prev
            results[-1] = (kind, f"{prev_text} / {seg}")
        else:
            results.append((False, seg))
    return results or [(True, cleaned)]


def _append_text(lesson: ParsedLesson, text: str) -> None:
    """Продолжение пары в следующей строке (напр. преподаватель отдельной строкой)."""
    extra = split_lesson_text(text)
    if not lesson.teacher and extra.teacher:
        lesson.teacher = extra.teacher
    if not lesson.room and extra.room:
        lesson.room = extra.room
    if extra.subject and extra.subject not in (lesson.subject or ""):
        lesson.subject = f"{lesson.subject} {extra.subject}".strip() if lesson.subject else extra.subject


def _find_date_in_grid(grid: list[list[str]], scan_rows: int = 12) -> date | None:
    for row in grid[:scan_rows]:
        for text in row:
            if not text:
                continue
            day = extract_date(text)
            if day:
                return day
    return None


# ------------------------------------------------------------------- основное
def _parse_sheet(ws, default_date: date | None) -> list[GroupSchedule]:
    grid = _sheet_grid(ws)
    if not grid:
        return []

    sheet_date = _find_date_in_grid(grid) or default_date
    course_match = COURSE_RE.search(ws.title or "")
    current_course = int(course_match.group(1)) if course_match else None

    result: dict[str, GroupSchedule] = {}
    columns: list[_GroupColumn] = []

    for row in grid:
        joined = " ".join(t for t in row if t)
        if not joined:
            continue

        course_match = COURSE_RE.search(joined)
        found = _find_groups_in_row(row)
        is_header = len(found) >= 2 or (not columns and len(found) == 1)

        if is_header:
            columns = _build_columns(found, len(row))
            if course_match:
                current_course = int(course_match.group(1))
            continue

        if course_match and not found:
            current_course = int(course_match.group(1))

        if not columns:
            continue

        number = _find_lesson_number(row, columns[0].start)

        for column in columns:
            texts = _cell_texts(row, column.start, column.end, column.group)
            if not texts:
                continue
            text = " ".join(texts)

            schedule = result.setdefault(
                column.group,
                GroupSchedule(
                    group=column.group, schedule_date=sheet_date, course=current_course
                ),
            )
            if schedule.course is None:
                schedule.course = current_course

            # Классный час стоит вне нумерованных пар (в строке только время).
            # Храним его с номером 0, чтобы не потерять и не склеить с другой парой.
            if number is None and "классный час" in text.lower():
                number = 0

            if number is None:
                if schedule.lessons:
                    _append_text(schedule.lessons[-1], text)
                continue

            existing = next((l for l in schedule.lessons if l.number == number), None)
            if existing:
                _append_text(existing, text)
            else:
                schedule.lessons.append(_make_lesson(number, text))

    return list(result.values())


def parse_excel_file(
    path: str | Path, default_date: date | None = None
) -> list[GroupSchedule]:
    """Главная функция парсера: файл -> список расписаний по группам."""
    path = Path(path)
    if path.suffix.lower() == ".xls":
        return _parse_xls_file(path, default_date)
    try:
        workbook = load_workbook(path, data_only=True)
    except Exception:
        logger.exception("Не удалось открыть Excel-файл %s", path)
        return []

    merged: dict[str, GroupSchedule] = {}
    try:
        for ws in workbook.worksheets:
            try:
                for schedule in _parse_sheet(ws, default_date):
                    if not schedule.lessons:
                        continue
                    key = schedule.group
                    if key in merged:
                        target = merged[key]
                        known = {l.number for l in target.lessons}
                        for lesson in schedule.lessons:
                            if lesson.number not in known:
                                target.lessons.append(lesson)
                        target.course = target.course or schedule.course
                        target.schedule_date = target.schedule_date or schedule.schedule_date
                    else:
                        merged[key] = schedule
            except Exception:
                logger.exception("Ошибка разбора листа '%s' в файле %s", ws.title, path.name)
    finally:
        workbook.close()

    logger.info("Файл %s: распознано групп — %s", path.name, len(merged))
    return list(merged.values())


def _parse_xls_file(path: Path, default_date: date | None) -> list[GroupSchedule]:
    """Разбирает старый бинарный Excel-файл факультета Чернышевского."""
    workbook = xlrd.open_workbook(path.as_posix(), on_demand=True)
    result: dict[str, GroupSchedule] = {}
    months = {
        "января": 1, "февраля": 2, "марта": 3, "апреля": 4,
        "мая": 5, "июня": 6, "июля": 7, "августа": 8,
        "сентября": 9, "октября": 10, "ноября": 11, "декабря": 12,
    }
    for sheet in workbook.sheets():
        day = default_date
        match = re.search(r"(\d{1,2})\s+([а-яё]+)", sheet.name.lower())
        if match and match.group(2) in months:
            day = date(datetime.now().year, months[match.group(2)], int(match.group(1)))

        # xlrd нумерует строки с нуля. В файлах Чернышевского два блока:
        # I/II курс — заголовок курса в строке 1, группы в строке 2;
        # III/IV курс — заголовок в строке 10, группы в строке 11.
        # Второй блок ищем по номеру курса (I-IV курс), а не по простому вхождению
        # подстроки «курс», иначе слова вроде «экскурсионных» в тексте пары (V пара)
        # ошибочно принимаются за начало блока старших курсов и обрезают 5-6 пары!
        second = next(
            (
                r for r in range(8, sheet.nrows)
                if any(ROMAN_COURSE_RE.search(_clean(str(sheet.cell_value(r, c)))) for c in range(sheet.ncols))
            ),
            None,
        )
        blocks = [{"course_row": 1, "header_row": 2, "lesson_rows": range(3, second or 10)}]
        if second is not None and second + 1 < sheet.nrows:
            blocks.append({
                "course_row": second,
                "header_row": second + 1,
                "lesson_rows": range(second + 2, sheet.nrows),
            })

        for block in blocks:
            c_row = block["course_row"]
            h_row = block["header_row"]
            if h_row >= sheet.nrows:
                continue

            course_cols: list[tuple[int, int]] = []
            for c in range(sheet.ncols):
                val = _clean(str(sheet.cell_value(c_row, c)))
                cm = ROMAN_COURSE_RE.search(val)
                if cm:
                    num = ROMAN_TO_NUM.get(cm.group(1).upper())
                    if num:
                        course_cols.append((c, num))
            course_cols.sort(key=lambda x: x[0])

            def get_course_for_col(col_idx: int) -> int:
                active_c = 1 if c_row == 1 else 3
                for c_pos, c_num in course_cols:
                    if col_idx >= c_pos:
                        active_c = c_num
                return active_c

            time_cols_for_block = []
            for c in range(sheet.ncols):
                has_num = False
                for r in block["lesson_rows"]:
                    val = _clean(str(sheet.cell_value(r, c)))
                    if parse_lesson_number(val) is not None:
                        has_num = True
                        break
                if has_num:
                    time_cols_for_block.append(c)

            def get_time_col(col_idx: int) -> int:
                best = 1
                for tc in time_cols_for_block:
                    if tc <= col_idx:
                        best = tc
                return best

            col_groups: list[tuple[int, str, int, int]] = []
            for c in range(sheet.ncols):
                raw_cell = str(sheet.cell_value(h_row, c))
                for line in raw_cell.splitlines():
                    grp = normalize_group(line)
                    if grp:
                        col_groups.append((c, grp, get_course_for_col(c), get_time_col(c)))

            for r in block["lesson_rows"]:
                for col_idx, group_name, course_num, time_col_idx in col_groups:
                    text = _clean(str(sheet.cell_value(r, col_idx)))
                    if not text:
                        continue

                    time_val = _clean(str(sheet.cell_value(r, time_col_idx)))
                    number = parse_lesson_number(time_val)

                    if number is None:
                        for tc in sorted(time_cols_for_block, key=lambda c: abs(c - col_idx)):
                            cand = parse_lesson_number(_clean(str(sheet.cell_value(r, tc))))
                            if cand is not None:
                                number = cand
                                break

                    schedule = result.setdefault(
                        group_name,
                        GroupSchedule(group=group_name, schedule_date=day, course=course_num)
                    )

                    for is_class_hour, part in _split_class_hour_cell(text):
                        if is_class_hour:
                            embedded = re.search(
                                r"с\s*(\d{1,2})[.:]\s*(\d{2})\s*классный", part, re.IGNORECASE
                            )
                            # Классный час середины дня (в строке нумерованной пары
                            # или с «с 14.10» в тексте) получает номер 90;
                            # утренний (отдельная строка без пары) — номер 0.
                            ch_number = (
                                MIDDAY_CLASS_HOUR
                                if (embedded or number is not None)
                                else 0
                            )
                            existing_ch = next(
                                (l for l in schedule.lessons if l.number == ch_number), None
                            )
                            if existing_ch is None:
                                lesson = _make_lesson(ch_number, part)
                                lesson.subject = (
                                    CLASS_HOUR_TIME_PREFIX_RE.sub("", lesson.subject or "").strip()
                                    or "Классный час"
                                )
                                schedule.lessons.append(lesson)
                        else:
                            if number is None:
                                continue
                            existing = next(
                                (l for l in schedule.lessons if l.number == number), None
                            )
                            if existing:
                                _append_text(existing, part)
                            else:
                                schedule.lessons.append(_make_lesson(number, part))

    workbook.release_resources()
    return list(result.values())


def find_group_schedule(
    path: str | Path, group: str, default_date: date | None = None
) -> GroupSchedule | None:
    """Первая цель ТЗ: расписание ОДНОЙ конкретной группы (для отладки)."""
    target = normalize_group(group) or group.upper()
    for schedule in parse_excel_file(path, default_date):
        if schedule.group == target:
            return schedule
    return None
