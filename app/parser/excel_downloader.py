"""ШАГ 13: скачивание Excel с проверками."""
from __future__ import annotations

import hashlib
import logging
from dataclasses import dataclass
from datetime import date
from io import BytesIO
from pathlib import Path

import httpx
from openpyxl import load_workbook
import xlrd

logger = logging.getLogger(__name__)

MIN_FILE_SIZE = 1024              # меньше 1 КБ — точно не расписание
MAX_FILE_SIZE = 25 * 1024 * 1024  # защита от мусора
HEADERS = {"User-Agent": "Mozilla/5.0 (compatible; PKPS-Schedule-Bot/1.0)"}


class DownloadError(Exception):
    """Любая проблема скачивания/проверки файла."""


@dataclass
class DownloadResult:
    path: Path
    size: int
    sha256: str
    sheet_names: list[str]


def build_file_name(schedule_date: date | None, file_type: str, url: str) -> str:
    day = schedule_date.isoformat() if schedule_date else "unknown-date"
    suffix = Path(url.split("?")[0]).suffix.lower() or ".xlsx"
    if suffix not in (".xlsx", ".xlsm", ".xls"):
        suffix = ".xlsx"
    return f"{day}_{file_type}{suffix}"


def _unique_path(directory: Path, file_name: str) -> Path:
    path = directory / file_name
    if not path.exists():
        return path
    stem, suffix = path.stem, path.suffix
    for i in range(2, 100):
        candidate = directory / f"{stem}_{i}{suffix}"
        if not candidate.exists():
            return candidate
    return directory / f"{stem}_last{suffix}"


async def download_excel(
    url: str, dest_dir: Path, file_name: str, timeout: float = 60.0
) -> DownloadResult:
    """Скачивает файл, проверяет его и сохраняет в data/excel."""
    dest_dir.mkdir(parents=True, exist_ok=True)

    try:
        async with httpx.AsyncClient(
            headers=HEADERS, timeout=timeout, follow_redirects=True
        ) as client:
            response = await client.get(url)
            response.raise_for_status()
            content = response.content
    except httpx.HTTPError as exc:
        raise DownloadError(f"Ошибка сети при скачивании: {exc}") from exc

    size = len(content)
    if size < MIN_FILE_SIZE:
        raise DownloadError(f"Файл слишком маленький ({size} байт)")
    if size > MAX_FILE_SIZE:
        raise DownloadError(f"Файл слишком большой ({size} байт)")
    is_xls = content.startswith(b"\xd0\xcf\x11\xe0")
    if not (content.startswith(b"PK") or is_xls):
        raise DownloadError("Файл не является Excel (.xlsx/.xls)")

    # проверяем, что файл реально открывается как Excel
    try:
        if is_xls:
            workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
            sheet_names = workbook.sheet_names()
            workbook.release_resources()
        else:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            sheet_names = list(workbook.sheetnames)
            workbook.close()
    except Exception as exc:  # повреждённый архив и т.п.
        raise DownloadError(f"Файл не открывается как Excel: {exc}") from exc

    path = _unique_path(dest_dir, file_name)
    try:
        path.write_bytes(content)
    except OSError as exc:
        raise DownloadError(f"Не удалось сохранить файл: {exc}") from exc

    digest = hashlib.sha256(content).hexdigest()
    logger.info("Скачан файл %s (%s КБ, листы: %s)", path.name, size // 1024, sheet_names)
    return DownloadResult(path=path, size=size, sha256=digest, sheet_names=sheet_names)
