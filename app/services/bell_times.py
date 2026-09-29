"""Звонки из исходного Excel: концы — лишь при согласии совпавших профилей,
иначе fallback; блоки листов/курсов не смешиваются; битый файл — пусто."""
from __future__ import annotations

import logging
import re
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

import xlrd
from openpyxl import load_workbook

from app.database.database import get_session
from app.database.models import ScheduleFile
from app.parser.excel_parser import ROMAN_TO_NUM
from app.utils import (
    LESSON_TIMES_CHERN_MONDAY_1, LESSON_TIMES_CHERN_MONDAY_24, LESSON_TIMES_CHERN_SATURDAY,
    LESSON_TIMES_CHERN_WEEKDAY_1, LESSON_TIMES_CHERN_WEEKDAY_24, LESSON_TIMES_MONDAY,
    LESSON_TIMES_SATURDAY, LESSON_TIMES_WEEKDAY, normalize_group,
)

logger = logging.getLogger(__name__)

# Номер + пробел + время (+ опциональный конец). Пробел обязателен, иначе «14:10» дало бы «1-ю пару в 04:10».
_BELL_CELL_RE = re.compile(
    r"^\s*(?P<number>[IVX]+|\d{1,2})\s+(?:пара|пар[аы]?|п\.)?\s*"
    r"(?P<hour>\d{1,2})[:.](?P<minute>\d{2})"
    r"(?:\s*[-–—]\s*(?P<end_hour>\d{1,2})[:.](?P<end_minute>\d{2}))?\s*$",
    re.IGNORECASE,
)
_PROFILES = (
    LESSON_TIMES_MONDAY, LESSON_TIMES_WEEKDAY, LESSON_TIMES_SATURDAY,
    LESSON_TIMES_CHERN_MONDAY_1, LESSON_TIMES_CHERN_MONDAY_24, LESSON_TIMES_CHERN_WEEKDAY_1,
    LESSON_TIMES_CHERN_WEEKDAY_24, LESSON_TIMES_CHERN_SATURDAY,
)
MIN_COLUMN_CELLS = 2  # колонка звонков: >= 2 разных номеров
MIN_BLOCK_LESSONS = 3  # короче — профиль не подбираем (одно 11:10 — не «по часу»)
MAX_GROUP_SCAN_ROWS = 40


@dataclass
class _BellBlock:
    sheet: str
    zone: tuple[int, int]  # строки шапки блока (поиск группы при расхождении)
    starts: dict[int, str]
    explicit: dict[int, str]


def _format_time(hour: str, minute: str) -> str | None:
    h, m = int(hour), int(minute)
    return f"{h:02d}:{m:02d}" if 0 <= h <= 23 and 0 <= m <= 59 else None


def _parse_bell_cell(value: object) -> tuple[int, str, str | None] | None:
    """(номер, начало, конец|None); чистое время/предметы/группы — None."""
    match = _BELL_CELL_RE.match(str(value or "").replace("l", "I").replace("L", "I").strip())
    if not match:
        return None
    raw = match.group("number")
    number = int(raw) if raw.isdigit() else ROMAN_TO_NUM.get(raw.upper())
    start = _format_time(match.group("hour"), match.group("minute"))
    raw_end = match.group("end_hour")
    end = _format_time(raw_end, match.group("end_minute")) if raw_end else None
    if number is None or not 1 <= number <= 8 or start is None or (raw_end and not end):
        return None
    return number, start, end


def _scan_column(items: list[tuple[object, int]]) -> list[tuple[int, str, str | None, int]] | None:
    """[(значение, строка)] -> вхождения звонков; конфликт номеров — None."""
    occurrences: list[tuple[int, str, str | None, int]] = []
    seen: dict[int, tuple[str, str | None]] = {}
    for value, row in items:
        parsed = _parse_bell_cell(value)
        if parsed is None:
            continue
        number, start, end = parsed
        if seen.get(number, (start, end)) != (start, end):
            return None
        seen.setdefault(number, (start, end))
        occurrences.append((number, start, end, row))
    return occurrences if len(seen) >= MIN_COLUMN_CELLS else None


def _make_block(sheet: str, zone: tuple[int, int],
                cells: list[tuple[int, str, str | None, int]]) -> _BellBlock | None:
    """Сливает ячейки; конфликт номеров (разные времена) — None."""
    merged: dict[int, tuple[str, str | None, int]] = {}
    for number, start, end, row in cells:
        if merged.get(number, (start, end, row))[:2] != (start, end):
            return None
        merged.setdefault(number, (start, end, row))
    if not merged:
        return None
    return _BellBlock(sheet, zone, {n: v[0] for n, v in merged.items()},
                      {n: v[1] for n, v in merged.items() if v[1]})


def _read_xlsx_blocks(path: Path) -> list[_BellBlock]:
    workbook = load_workbook(path, data_only=True)
    blocks: list[_BellBlock] = []
    try:
        for sheet in workbook.worksheets:
            if "курс" not in (sheet.title or "").lower():
                continue
            cells = [item for column in sheet.iter_cols()
                     for item in (_scan_column([(c.value, c.row) for c in column]) or [])]
            if not cells:
                continue
            first = min(row for _, _, _, row in cells)
            block = _make_block(sheet.title or "", (1, first - 1), cells)
            if block is not None:
                blocks.append(block)
    finally:
        workbook.close()
    return blocks


def _read_xls_blocks(path: Path) -> list[_BellBlock]:
    workbook = xlrd.open_workbook(path.as_posix())
    blocks: list[_BellBlock] = []
    try:
        for sheet in workbook.sheets():
            columns = [
                found for col in range(sheet.ncols)
                if (found := _scan_column(
                    [(sheet.cell_value(row, col), row + 1)
                     for row in range(sheet.nrows)])) is not None
            ]
            rows = sorted({row for column in columns for _, _, _, row in column})
            if not rows:
                continue
            segments = [[rows[0]]]
            for row in rows[1:]:
                if row - segments[-1][-1] > 2:  # разрыв > 2 (шапка) — новый блок курсов
                    segments.append([])
                segments[-1].append(row)
            previous_end = 0
            for segment in segments:
                in_segment = set(segment)
                block = _make_block(
                    sheet.name, (previous_end + 1, segment[0] - 1),
                    [(n, s, e, r) for column in columns for n, s, e, r in column if r in in_segment],
                )
                if block is not None:
                    blocks.append(block)
                previous_end = segment[-1]
    finally:
        workbook.release_resources()
    return blocks


def _find_group_locations(path: Path, group: str) -> list[tuple[str, int]]:
    """(лист, строка) заголовков с группой (построчно, шапка листа)."""
    hits: list[tuple[str, int]] = []

    def scan(value: object, sheet: str, row: int) -> None:
        hits.extend((sheet, row) for line in str(value or "").splitlines()
                    if normalize_group(line) == group)

    if path.suffix.lower() == ".xls":
        workbook = xlrd.open_workbook(path.as_posix())
        try:
            for sheet in workbook.sheets():
                for row in range(min(sheet.nrows, MAX_GROUP_SCAN_ROWS)):
                    for col in range(sheet.ncols):
                        scan(sheet.cell_value(row, col), sheet.name, row + 1)
        finally:
            workbook.release_resources()
    else:
        workbook = load_workbook(path, data_only=True, read_only=True)
        try:
            for sheet in workbook.worksheets:
                for index, row in enumerate(sheet.iter_rows(max_row=MAX_GROUP_SCAN_ROWS), 1):
                    for cell in row:
                        scan(cell.value, sheet.title or "", index)
        finally:
            workbook.close()
    return hits


def _resolve_block(blocks: list[_BellBlock], path: Path, group: str) -> _BellBlock | None:
    """Одинаковые блоки — любой; расходятся — блок группы; не нашли — None."""
    unique = {frozenset(block.starts.items()): block for block in blocks}
    if len(unique) == 1:
        return next(iter(unique.values()))
    if not group:
        return None
    locations = _find_group_locations(path, group)
    matched = [block for block in blocks if any(
        sheet == block.sheet and block.zone[0] <= row <= block.zone[1]
        for sheet, row in locations)]
    return matched[0] if len(matched) == 1 else None


def _agreeing_ends(starts: dict[int, str]) -> dict[int, str | None]:
    """Конец — только если все совпавшие с блоком профили согласны."""
    candidates = [t for t in _PROFILES if all(n in t and t[n][0] == s for n, s in starts.items())]
    ends: dict[int, str | None] = {}
    for number in starts:
        values = {table[number][1] for table in candidates if number in table}
        ends[number] = next(iter(values)) if len(values) == 1 else None
    return ends


@lru_cache(maxsize=256)
def _lesson_ranges(path_str: str, modified_ns: int, group: str) -> dict[int, tuple[str, str]]:
    path = Path(path_str)
    try:
        blocks = (_read_xls_blocks(path) if path.suffix.lower() == ".xls"
                  else _read_xlsx_blocks(path))
        block = _resolve_block(blocks, path, group) if blocks else None
    except FileNotFoundError:
        return {}
    except Exception:
        logger.warning("Звонки из %s не прочитаны", path.name, exc_info=True)
        return {}
    if block is None:
        return {}
    agreed = _agreeing_ends(block.starts) if len(block.starts) >= MIN_BLOCK_LESSONS else {}
    return {n: (s, e) for n, s in block.starts.items()
            if (e := block.explicit.get(n) or agreed.get(n))}


def get_lesson_time_range(file_id: int | None, group: str | None, lesson_number: int) -> tuple[str, str] | None:
    """(начало, конец) пары из файла, если однозначно; иначе None (fallback вызывающего)."""
    if not file_id:
        return None
    with get_session() as session:
        source = session.get(ScheduleFile, file_id)
        local_path = source.local_path if source else None
    if not local_path:
        return None
    path = Path(local_path)
    if not path.is_file():
        return None
    return _lesson_ranges(str(path.resolve()), path.stat().st_mtime_ns, group or "").get(lesson_number)
